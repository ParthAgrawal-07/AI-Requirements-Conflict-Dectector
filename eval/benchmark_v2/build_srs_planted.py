"""Build srs_planted.docx from srs_planted.md (Abhyudit — benchmark v2, Sprint 1).

The Markdown file is the editable source; the .docx is what people upload. Run from ``backend/``
(python-docx comes with the backend requirements)::

    python ../eval/benchmark_v2/build_srs_planted.py

Mapping: ``#``/``##``/``###`` -> Word Title/Heading 1/Heading 2 styles, consecutive ``| a | b |``
lines -> a Word table (first row is the header), every other non-empty line -> a paragraph.
The output is deterministic, so a rebuild without source changes gives the same text.
"""

from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "srs_planted.md"
TARGET = HERE / "srs_planted.docx"


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def build(source: Path = SOURCE, target: Path = TARGET) -> Path:
    document = docx.Document()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        table.style = "Table Grid"
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        table_rows.clear()

    for line in source.read_text(encoding="utf-8").splitlines():
        if line.startswith("|"):
            table_rows.append(_cells(line))
            continue
        flush_table()
        if not line.strip():
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=0)
        else:
            document.add_paragraph(line)
    flush_table()

    document.core_properties.title = "SecureDocs Portal SRS (benchmark v2)"
    document.save(target)
    return target


if __name__ == "__main__":
    print(f"wrote {build()}")
