"""Tests for DOCXParser."""

import docx
import pytest

from app.ingestion.errors import DocumentParseError
from app.ingestion.models import RawSegment
from app.ingestion.parsers.docx_parser import DOCXParser


class TestDOCXParser:
    """Tests for the DOCX parser."""

    def test_returns_list_of_raw_segments(self, sample_docx):
        parser = DOCXParser()
        segments = parser.parse(sample_docx)

        assert isinstance(segments, list)
        assert len(segments) > 0
        assert all(isinstance(s, RawSegment) for s in segments)

    def test_extracts_text_content(self, sample_docx):
        parser = DOCXParser()
        segments = parser.parse(sample_docx)

        assert segments[0].text  # non-empty
        assert "REQ-003" in segments[0].text

    def test_preserves_paragraph_index(self, sample_docx):
        parser = DOCXParser()
        segments = parser.parse(sample_docx)

        assert segments[0].page_or_paragraph == 1  # 1-based

    def test_preserves_source_document(self, sample_docx):
        parser = DOCXParser()
        segments = parser.parse(sample_docx)

        assert segments[0].source_document == "sample.docx"

    def test_skips_empty_paragraphs(self, sample_docx_multi):
        parser = DOCXParser()
        segments = parser.parse(sample_docx_multi)

        # 4 paragraphs, but 1 is empty → 3 segments
        assert len(segments) == 3
        texts = [s.text for s in segments]
        assert any("FR-1" in t for t in texts)
        assert any("NFR-02" in t for t in texts)
        # Ensure order_index is sequential
        assert [s.order_index for s in segments] == [0, 1, 2]

    def test_rejects_doc_format(self, tmp_path):
        """Explicitly test that .doc is rejected with a helpful message."""
        fake_file = tmp_path / "legacy.doc"
        fake_file.write_bytes(b"\x00" * 10)

        parser = DOCXParser()
        with pytest.raises(ValueError, match="not supported"):
            parser.parse(str(fake_file))

    def test_rejects_txt_format(self, tmp_path):
        fake_file = tmp_path / "notes.txt"
        fake_file.write_text("hello")

        parser = DOCXParser()
        with pytest.raises(ValueError, match="Unsupported file format"):
            parser.parse(str(fake_file))

    def test_rejects_nonexistent_file(self):
        parser = DOCXParser()
        with pytest.raises(FileNotFoundError, match="File not found"):
            parser.parse("/nonexistent/path/document.docx")

    def test_extracts_table_rows(self, sample_docx_structured):
        """Requirements inside tables must not be dropped."""
        segments = DOCXParser().parse(sample_docx_structured)

        texts = [s.text for s in segments]
        assert "REQ-010 | The system shall encrypt data at rest." in texts
        assert "REQ-011 | The system shall store data unencrypted." in texts

    def test_headings_become_section_heading(self, sample_docx_structured):
        """Headings are attached as context, not emitted as segments."""
        segments = DOCXParser().parse(sample_docx_structured)

        assert all("Functional Requirements" not in s.text for s in segments)
        assert all(s.section_heading == "3. Functional Requirements" for s in segments)
        assert all(s.location_type == "paragraph" for s in segments)

    def test_corrupted_docx_raises_parse_error(self, tmp_path):
        bad = tmp_path / "broken.docx"
        bad.write_bytes(b"not a zip archive")

        with pytest.raises(DocumentParseError, match="corrupted"):
            DOCXParser().parse(str(bad))

    def test_skips_table_of_contents_paragraphs(self, tmp_path):
        path = tmp_path / "toc.docx"
        doc = docx.Document()
        styles = doc.styles
        styles.add_style("TOC 1", 1)  # WD_STYLE_TYPE.PARAGRAPH
        doc.add_paragraph("1 Introduction\t2", style="TOC 1")
        doc.add_paragraph("REQ-1: The system shall work.")
        doc.save(str(path))

        assert [s.text for s in DOCXParser().parse(str(path))] == ["REQ-1: The system shall work."]

    def test_marks_list_paragraphs(self, realistic_srs_docx):
        segments = DOCXParser().parse(realistic_srs_docx)

        assert [s.text for s in segments if s.is_list_item] == ["PDF", "CSV"]
