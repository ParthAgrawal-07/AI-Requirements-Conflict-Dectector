"""
Shared pytest fixtures for ingestion tests.

Creates small deterministic PDF and DOCX test files on the fly
so tests don't depend on external sample documents.
"""

import docx
import pymupdf  # PyMuPDF
import pytest


@pytest.fixture
def sample_pdf(tmp_path):
    """Create a minimal PDF with two requirement lines on page 1."""
    path = tmp_path / "sample.pdf"

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "REQ-001: The system shall allow user login.\nREQ-002: The system shall support logout.",
    )
    doc.save(str(path))
    doc.close()

    return str(path)


@pytest.fixture
def sample_pdf_multipage(tmp_path):
    """Create a PDF with content on page 1, empty page 2, content on page 3."""
    path = tmp_path / "multipage.pdf"

    doc = pymupdf.open()

    # Page 1 — has content
    page1 = doc.new_page()
    page1.insert_text((50, 50), "FR-1: Functional requirement one.")

    # Page 2 — intentionally empty
    doc.new_page()

    # Page 3 — has content
    page3 = doc.new_page()
    page3.insert_text((50, 50), "NFR-01: Non-functional requirement.")

    doc.save(str(path))
    doc.close()

    return str(path)


@pytest.fixture
def sample_docx(tmp_path):
    """Create a minimal DOCX with one requirement paragraph."""
    path = tmp_path / "sample.docx"

    doc = docx.Document()
    doc.add_paragraph("REQ-003: The system shall export reports.")
    doc.save(str(path))

    return str(path)


@pytest.fixture
def sample_docx_multi(tmp_path):
    """Create a DOCX with mixed content: requirements, empty paras, and plain text."""
    path = tmp_path / "multi.docx"

    doc = docx.Document()
    doc.add_paragraph("FR-1: First functional requirement.")
    doc.add_paragraph("")  # empty paragraph — should be skipped
    doc.add_paragraph("This is an introductory paragraph with no ID.")
    doc.add_paragraph("NFR-02: Performance requirement.")
    doc.save(str(path))

    return str(path)


@pytest.fixture
def sample_docx_structured(tmp_path):
    """Create a DOCX with a heading and a requirements table (ID | Description)."""
    path = tmp_path / "structured.docx"

    doc = docx.Document()
    doc.add_heading("3. Functional Requirements", level=1)
    doc.add_paragraph("REQ-001: The system shall allow login.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "REQ-010"
    table.cell(0, 1).text = "The system shall encrypt data at rest."
    table.cell(1, 0).text = "REQ-011"
    table.cell(1, 1).text = "The system shall store data unencrypted."
    doc.save(str(path))

    return str(path)


def _write_pages(path, pages, *, header=None, footer=None):
    """Write a PDF where each page is a list of (text, fontsize) paragraphs, wrapped to width."""
    doc = pymupdf.open()
    width, height = 595, 842
    for number, paragraphs in enumerate(pages, start=1):
        page = doc.new_page(width=width, height=height)
        if header:
            page.insert_text((50, 30), header, fontsize=8)
        if footer:
            page.insert_text((270, 820), footer.format(n=number, total=len(pages)), fontsize=8)
        y = 90
        for text, size in paragraphs:
            box = pymupdf.Rect(50, y, width - 50, y + 200)
            unused = page.insert_textbox(box, text, fontsize=size)
            y += (200 - unused) + 12
    doc.save(str(path))
    doc.close()


@pytest.fixture
def realistic_srs_pdf(tmp_path):
    """Three-page SRS: title + TOC, running header/footer, wrapped lines, a requirement
    split across a page break, a bullet list, ID-less statements and a duplicate ID."""
    path = tmp_path / "library_srs.pdf"
    _write_pages(
        path,
        [
            [
                ("Software Requirements Specification", 16),
                ("Online Library System", 12),
                ("Table of Contents", 12),
                ("1 Introduction ........................................ 2", 10),
                ("2 Functional Requirements ............................. 2", 10),
                ("3 Non-Functional Requirements ......................... 3", 10),
            ],
            [
                ("1 Introduction", 13),
                (
                    "This document describes the requirements for the Online Library System "
                    "used by members and librarians of the city library.",
                    10,
                ),
                ("2 Functional Requirements", 13),
                (
                    "REQ-001: The system shall allow a registered member to search the catalogue "
                    "by title, author, ISBN or subject, and shall return results within two "
                    "seconds for up to 10,000 records.",
                    10,
                ),
                ("REQ-002: The system shall support the following notification channels:", 10),
                ("- Email", 10),
                ("- SMS", 10),
                (
                    "REQ-003: A member must not borrow more than five items at the same time. "
                    "This limit applies across all branches.",
                    10,
                ),
                (
                    "Librarians shall be able to override the borrowing limit for a single "
                    "member. The override shall be logged.",
                    10,
                ),
                (
                    "REQ-004: The system shall send a reminder to the member three days before "
                    "an item is due and again on the day the item becomes overdue, unless the "
                    "member has",
                    10,
                ),
            ],
            [
                ("disabled reminders in their profile settings.", 10),
                ("3 Non-Functional Requirements", 13),
                ("NFR-01: The system should be available 99.5% of the time, measured monthly.", 10),
                ("The user interface will be easy to use.", 10),
                ("REQ-003: The catalogue shall be updated nightly.", 10),
            ],
        ],
        header="Online Library System - SRS v1.0",
        footer="Page {n} of {total}",
    )
    return str(path)


@pytest.fixture
def realistic_srs_docx(tmp_path):
    """DOCX with styled headings, a glossary, Word bullet list, and a table with a header row."""
    path = tmp_path / "library_srs.docx"

    doc = docx.Document()
    doc.add_heading("Online Library System SRS", 0)
    doc.add_heading("1 Introduction", 1)
    doc.add_paragraph("This SRS describes the Online Library System.")
    doc.add_heading("1.1 Glossary", 2)
    doc.add_paragraph("Member: a registered library user.")
    doc.add_heading("2 Functional Requirements", 1)
    doc.add_paragraph(
        "FR-1: The system shall let members reserve items. Reservations expire after 48 hours."
    )
    doc.add_paragraph("FR-2: The system shall export reports in these formats:")
    doc.add_paragraph("PDF", style="List Bullet")
    doc.add_paragraph("CSV", style="List Bullet")
    doc.add_paragraph("The system shall lock an account after five failed login attempts.")
    doc.add_heading("2.1 Fines", 2)
    doc.add_paragraph("Fines are calculated daily.")
    table = doc.add_table(rows=3, cols=3)
    rows = [
        ("ID", "Requirement", "Priority"),
        ("FR-10", "The system shall calculate fines daily.", "High"),
        ("FR-11", "Fines must not exceed the item price.", "Medium"),
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_heading("3 Non-Functional Requirements", 1)
    doc.add_paragraph("NFR-1: Pages should load within 2 seconds.")
    doc.save(str(path))

    return str(path)
