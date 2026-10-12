"""Integration tests for storing extraction results (app.ingestion.service, US-02)."""

import shutil
import uuid
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from pathlib import Path

import docx
import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.ingestion import service, tasks
from app.ingestion.service import extract_document
from app.models import (
    Document,
    DocumentFormat,
    DocumentStatus,
    Requirement,
    RequirementVersion,
    User,
    Workspace,
)


@pytest.fixture
def make_document(
    db_session: Session, workspace: Workspace, user: User, tmp_path: Path
) -> Callable[..., Document]:
    """Store `source` under a server-side name (as the upload endpoint would) and add a row."""

    def factory(source: str | Path | None, *, filename: str = "Library SRS.pdf") -> Document:
        suffix = Path(filename).suffix
        stored = tmp_path / f"{uuid.uuid4().hex}{suffix}"
        if source is not None:
            shutil.copy(source, stored)
        document = Document(
            workspace_id=workspace.id,
            uploaded_by=user.id,
            filename=filename,
            format=DocumentFormat(suffix.lstrip(".")),
            size_bytes=stored.stat().st_size if source is not None else 0,
            sha256="0" * 64,
            storage_path=str(stored),
        )
        db_session.add(document)
        db_session.commit()
        return document

    return factory


def _requirements(session: Session, document: Document) -> list[Requirement]:
    return list(
        session.scalars(
            select(Requirement)
            .where(Requirement.document_id == document.id)
            .order_by(Requirement.order_index)
        )
    )


def test_stores_requirements_and_marks_done(db_session, make_document, sample_pdf):
    document = make_document(sample_pdf)

    extract_document(db_session, document.id)

    assert document.status == DocumentStatus.DONE
    assert document.progress_percent == 100
    assert document.error_message is None
    reqs = _requirements(db_session, document)
    assert [r.source_req_id for r in reqs] == ["REQ-001", "REQ-002"]  # preserved (DR-03)
    assert reqs[0].body == "REQ-001: The system shall allow user login."
    assert reqs[0].location == "page 1"
    assert [r.order_index for r in reqs] == [0, 1]
    assert not any(r.needs_review for r in reqs)


def test_creates_initial_version_snapshot(db_session, make_document, sample_pdf):
    document = make_document(sample_pdf)

    extract_document(db_session, document.id)

    versions = db_session.scalars(
        select(RequirementVersion).join(Requirement).where(Requirement.document_id == document.id)
    ).all()
    assert len(versions) == 2
    assert all(v.version_no == 1 and v.changed_by is None for v in versions)
    assert {v.body_snapshot for v in versions} == {
        r.body for r in _requirements(db_session, document)
    }


def test_persists_review_flags_and_section(db_session, make_document, tmp_path):
    path = tmp_path / "dupes.docx"
    doc = docx.Document()
    doc.add_heading("2 Functional Requirements", level=1)
    doc.add_paragraph("REQ-001: The system shall allow login.")
    doc.add_paragraph("REQ-001: The system shall allow logout.")
    doc.add_paragraph("The system shall lock idle sessions.")
    doc.save(str(path))
    document = make_document(path, filename="dupes.docx")

    extract_document(db_session, document.id)

    reqs = _requirements(db_session, document)
    assert [r.needs_review for r in reqs] == [True, True, False]
    assert "Duplicate requirement ID REQ-001" in (reqs[0].review_reason or "")
    assert reqs[2].source_req_id is None
    assert reqs[2].review_reason is None
    assert all(r.section_heading == "2 Functional Requirements" for r in reqs)


def test_corrupted_file_fails_with_uploaded_filename(db_session, make_document, tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"definitely not a pdf")
    document = make_document(bad, filename="Library SRS.pdf")

    extract_document(db_session, document.id)

    assert document.status == DocumentStatus.FAILED
    assert "'Library SRS.pdf' is not a valid PDF" in (document.error_message or "")
    assert Path(document.storage_path).name not in (document.error_message or "")
    assert _requirements(db_session, document) == []


def test_missing_file_fails_without_leaking_storage_path(db_session, make_document):
    document = make_document(None)

    extract_document(db_session, document.id)

    assert document.status == DocumentStatus.FAILED
    assert document.error_message == "The uploaded file could not be found on the server."


def test_document_without_requirements_fails(db_session, make_document, tmp_path):
    path = tmp_path / "intro.docx"
    doc = docx.Document()
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph("This document describes the library system.")
    doc.save(str(path))
    document = make_document(path, filename="intro.docx")

    extract_document(db_session, document.id)

    assert document.status == DocumentStatus.FAILED
    assert "No requirement statements were found" in (document.error_message or "")


def test_rerun_is_idempotent(db_session, make_document, sample_pdf):
    document = make_document(sample_pdf)

    extract_document(db_session, document.id)
    extract_document(db_session, document.id)

    assert len(_requirements(db_session, document)) == 2


def test_replaces_rows_left_by_an_interrupted_run(db_session, make_document, sample_pdf):
    document = make_document(sample_pdf)
    db_session.add(Requirement(document_id=document.id, body="half-written leftover"))
    db_session.commit()

    extract_document(db_session, document.id)

    bodies = [r.body for r in _requirements(db_session, document)]
    assert "half-written leftover" not in bodies
    assert len(bodies) == 2


def test_unknown_document_returns_none(db_session):
    assert extract_document(db_session, uuid.uuid4()) is None


def test_unexpected_error_marks_failed_and_reraises(
    db_session, make_document, sample_pdf, monkeypatch
):
    document = make_document(sample_pdf)

    def boom(_path: str) -> None:
        raise RuntimeError("parser bug")

    monkeypatch.setattr(service, "ingest_document", boom)

    with pytest.raises(RuntimeError, match="parser bug"):
        extract_document(db_session, document.id)

    assert document.status == DocumentStatus.FAILED
    assert document.error_message == "Unexpected error while extracting requirements."


def test_celery_task_runs_extraction(db_session, make_document, sample_pdf, monkeypatch):
    document = make_document(sample_pdf)

    @contextmanager
    def test_scope() -> Iterator[Session]:
        yield db_session

    monkeypatch.setattr(tasks, "session_scope", test_scope)

    assert tasks.extract_document_task.run(str(document.id)) == "done"
    count = db_session.scalar(
        select(func.count()).select_from(Requirement).where(Requirement.document_id == document.id)
    )
    assert count == 2
